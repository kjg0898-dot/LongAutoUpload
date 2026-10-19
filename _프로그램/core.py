"""파일·메모장·엑셀·예약 계산. YouTube에 연결하지 않는 부분만 모아둔다."""
import difflib
import json
import re
import shutil
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

PROG = Path(__file__).resolve().parent
ROOT = PROG.parent
VIDEO_DIR = ROOT / "영상"
DONE_DIR = ROOT / "완료"
EXCEL = ROOT / "업로드목록.xlsx"
CONFIG = PROG / "설정.json"
KST = timezone(timedelta(hours=9))  # 한국은 서머타임이 없어서 고정 시차로 충분

VIDEO_EXT = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".wmv", ".flv"}
IMAGE_EXT = {".jpg", ".jpeg", ".png"}
THUMB_MAX = 2 * 1024 * 1024  # YouTube 썸네일 한도

KEEP = "건들지 않음"
MODES = ["맨 앞", "맨 뒤", "기준 글자 앞", "기준 글자 뒤", "표시 문자 바꾸기"]
CATEGORIES = {"영화/애니메이션": "1", "자동차/교통": "2", "음악": "10", "애완동물/동물": "15", "스포츠": "17",
              "여행/이벤트": "19", "게임": "20", "인물/블로그": "22", "코미디": "23", "엔터테인먼트": "24",
              "뉴스/정치": "25", "노하우/스타일": "26", "교육": "27", "과학기술": "28", "비영리/사회운동": "29"}
LANGUAGES = {"한국어": "ko", "영어": "en", "일본어": "ja", "중국어": "zh", "스페인어": "es"}

DEFAULT_CHANNEL = {
    "name": "", "folder": "", "ypp": False, "members": False, "slots": "월 18:00, 수 18:00, 금 18:00",
    # 초안 제목 틀이 "[채널명]  l 역사 인물 l ..." 꼴이라 "]" 바로 뒤에 넣는다
    "title_mode": "기준 글자 뒤", "title_marker": "]",
    "desc_mode": "맨 앞", "desc_marker": "",
    "tags_mode": "뒤에 추가", "playlist": "",
    "category": KEEP, "language": KEEP, "synthetic": KEEP, "kids": KEEP,
    "embeddable": KEEP, "stats": KEEP, "license": KEEP,
    "channel_id": "", "channel_title": "",
}
DEFAULT_CONFIG = {"channels": [], "max_fails": 3, "mode": "한 번에 전체"}

COLUMNS = ["채널", "파일", "제목", "제목후보", "설명", "태그", "고정댓글", "메모장", "썸네일",
           "회원우선", "예약일시", "상태", "오류", "영상ID"]


# ---------- 설정 ----------
def new_channel(name):
    return {**DEFAULT_CHANNEL, "key": uuid.uuid4().hex[:8], "name": name, "folder": name}


def load_config():
    cfg = {**DEFAULT_CONFIG, **(json.loads(CONFIG.read_text("utf-8")) if CONFIG.exists() else {})}
    cfg["channels"] = [{**DEFAULT_CHANNEL, **ch} for ch in cfg["channels"]]
    return cfg


def save_config(cfg):
    names = [ch["name"].strip() for ch in cfg["channels"]]
    if "" in names or len(set(names)) != len(names):
        raise ValueError("채널 이름이 비어 있거나 겹칩니다.")
    CONFIG.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), "utf-8")
    VIDEO_DIR.mkdir(exist_ok=True)


def overrides(ch):
    """'건들지 않음'이 아닌 설정만 골라 (snippet, status)로 돌려준다."""
    snippet, status = {}, {}
    if ch["category"] != KEEP:
        snippet["categoryId"] = CATEGORIES[ch["category"]]
    if ch["language"] != KEEP:
        snippet["defaultLanguage"] = snippet["defaultAudioLanguage"] = LANGUAGES[ch["language"]]
    yes = {"예": True, "아니요": False, "허용": True, "허용 안 함": False, "공개": True, "비공개": False}
    for key, field in (("synthetic", "containsSyntheticMedia"), ("kids", "selfDeclaredMadeForKids"),
                       ("embeddable", "embeddable"), ("stats", "publicStatsViewable")):
        if ch[key] != KEEP:
            status[field] = yes[ch[key]]
    if ch["license"] != KEEP:
        status["license"] = "creativeCommon" if "커먼즈" in ch["license"] else "youtube"
    return snippet, status


# ---------- 메모장 ----------
_SECTIONS = {"제목": "title", "타이틀": "title", "title": "title",
             "설명": "desc", "설명글": "desc", "본문": "desc", "description": "desc",
             "태그": "tags", "해시태그": "tags", "tag": "tags", "tags": "tags",
             "고정댓글": "comment", "댓글": "comment", "comment": "comment"}
# shortcut: 메모장 실제 형식을 보기 전이라 "제목:", "[설명]", "■ 태그" 같은 흔한 머리말을 추측한다
_HEADER = re.compile(r"^[\W_]*(제목|타이틀|설명글|설명|본문|해시태그|태그|고정\s*댓글|댓글|title|description|tags?|comment)"
                     r"\s*[\]\)】>》]*\s*(?:[:：]\s*(.*))?$", re.I)


def read_text(path):
    data = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp949"):  # 메모장은 UTF-8 또는 옛 한글 인코딩으로 저장된다
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def parse_memo(text):
    out, key, buf = {}, None, []
    for line in text.splitlines() + [None]:
        m = _HEADER.match(line.strip()) if line is not None else None
        if line is None or m:
            if key:
                out[key] = "\n".join(buf).strip()
            if m:
                key, buf = _SECTIONS[re.sub(r"\s", "", m.group(1)).lower()], [m.group(2) or ""]
        elif key:
            buf.append(line)
    return out


# SEO 메모장: 영상 여러 개가 "[날짜] 행 N | 순번 N | 제목: 원본 제목" 블록으로 이어지고, 블록 안은 --- 로 구역이 나뉜다
_SEO_HEAD = re.compile(r"^\[\d{4}-\d\d-\d\d[^\]]*\]\s*행\s*\d+\s*\|\s*순번\s*(\d+)\s*\|\s*제목:\s*(.*)$")


def parse_seo(text):
    blocks = []
    for line in text.splitlines():
        m = _SEO_HEAD.match(line.strip())
        if m:
            blocks.append((int(m[1]), m[2].strip(), []))
        elif blocks and not re.fullmatch(r"=+", line.strip()):
            blocks[-1][2].append(line)
    out = []
    for num, orig, lines in blocks:
        sections = {}
        for chunk in re.split(r"^---\s*$", "\n".join(lines), flags=re.M):
            head, _, body = chunk.strip().partition("\n")
            sections[head.strip()] = body.strip()

        def find(word):
            return next((v for k, v in sections.items() if k.startswith(word)), "")
        titles = re.findall(r"^\s*\d+\.\s*(.+?)\s*$", find("제목"), re.M)
        rec = re.search(r"추천\s*1순위:\s*(\d+)", find("제목"))
        pick = int(rec[1]) - 1 if rec and 0 < int(rec[1]) <= len(titles) else 0
        desc = find("설명")
        out.append({"num": num, "orig": orig, "titles": titles, "title": titles[pick] if titles else orig,
                    "desc": desc, "tags": "", "comment": find("고정")})  # 해시태그는 설명에만
    return out


def load_entries(memo_paths):
    """메모장들 → 영상 하나당 항목 하나. SEO 형식이면 한 파일에 여러 항목."""
    entries = []
    for path in memo_paths:
        text = read_text(path)
        seo = parse_seo(text)
        if seo:
            entries += [e | {"src": path} for e in seo]
        elif info := parse_memo(text):
            entries.append({"num": None, "orig": "", "titles": []} | info | {"src": path})
    return entries


def parse_tags(text):
    return [t.strip() for t in re.split(r"[,\n#]+", text or "") if t.strip()]


# ---------- 합치기 ----------
def merge(template, content, mode, marker="", sep=" "):
    """초안의 틀(template)은 남기고 content를 정해진 자리에 넣는다."""
    template, content = (template or "").strip(), (content or "").strip()
    if not content or content in template:  # 다시 실행해도 두 번 들어가지 않게
        return template
    if not template:
        return content
    if mode == "맨 앞":
        return content + sep + template
    if mode == "맨 뒤":
        return template + sep + content
    i = template.find(marker) if marker else -1
    if i < 0:
        raise ValueError(f"초안 틀에서 기준 글자 '{marker}'를 못 찾았습니다. 설정에서 맞춰주세요.")
    if mode == "표시 문자 바꾸기":
        return template[:i] + content + template[i + len(marker):]
    j = i if mode == "기준 글자 앞" else i + len(marker)
    return f"{template[:j].rstrip()}{sep}{content}{sep}{template[j:].lstrip()}".strip()


def tags_len(tags):
    # YouTube 계산법: 띄어쓰기가 있는 태그는 따옴표 2자, 태그 사이 쉼표 1자
    return sum(len(t) + (2 if " " in t else 0) for t in tags) + max(len(tags) - 1, 0)


def merge_tags(existing, new, mode):
    out, seen = [], set()
    for t in (new + existing if mode == "앞에 추가" else existing + new):
        t = t.strip().lstrip("#").strip()
        if t and t.lower() not in seen:
            seen.add(t.lower())
            out.append(t)
    while tags_len(out) > 500:
        out.pop()
    return out


def check_text(title, desc):
    if not title:
        raise ValueError("제목이 비어 있습니다.")
    if len(title) > 100:
        raise ValueError(f"제목이 100자를 넘습니다 ({len(title)}자): {title}")
    if len(desc.encode("utf-8")) > 5000:
        raise ValueError("설명이 YouTube 한도(5000바이트)를 넘습니다.")
    if re.search(r"[<>]", title + desc):
        raise ValueError("제목·설명에 < 또는 > 기호는 쓸 수 없습니다.")


# ---------- 파일 찾기 ----------
def natural_key(p):
    return [int(s) if s.isdigit() else s.lower() for s in re.split(r"(\d+)", str(p))]


def _norm(stem):
    stem = re.sub(r"썸네일|thumbnail|thumb|메모|설명|정보|info|최종|final", "", stem.lower())
    return re.sub(r"[\W_]", "", stem)


def _sim(a, b):
    return difflib.SequenceMatcher(None, _norm(a), _norm(b)).ratio() if _norm(a) and _norm(b) else 0


def _nums(path):
    """파일 이름 속 번호 (예: 숨겨진동양사이야기_6_내보내기 → {6}). 날짜 같은 큰 숫자는 뺀다."""
    return {int(n) for n in re.findall(r"\d+", Path(path).stem) if int(n) < 1000}


def _find_pair(video, candidates, videos):
    same_dir = [c for c in candidates if c.parent == video.parent]
    for c in same_dir:
        if c.stem.lower() == video.stem.lower():
            return c
    if len(same_dir) == 1 and sum(v.parent == video.parent for v in videos) == 1:
        return same_dir[0]  # 영상마다 폴더가 따로 있는 경우
    best = max(candidates, key=lambda c: _sim(c.stem, video.stem), default=None)
    return best if best and _sim(best.stem, video.stem) >= 0.6 else None


def _by_number(num, items, key):
    hits = [i for i in items if key(i) == {num}]
    return hits[0] if num is not None and len(hits) == 1 else None


def _match_entry(video, entries, videos):
    simple = [e for e in entries if e["num"] is None]
    src = _find_pair(video, [e["src"] for e in simple], videos)
    if src:
        return next(e for e in simple if e["src"] == src)
    nums = _nums(video)
    if hit := _by_number(next(iter(nums)) if len(nums) == 1 else None, entries, lambda e: {e["num"]}):
        return hit
    score = lambda e: max(_sim(video.stem, t) for t in [e["orig"], e["title"], *e["titles"]])  # noqa: E731
    best = max(entries, key=score, default=None)
    return best if best and score(best) >= 0.6 else None


def _match_thumb(video, entry, images, videos):
    same = [i for i in images if i.stem.lower() == video.stem.lower()]
    if same:
        return same[0]
    nums = _nums(video)
    num = entry["num"] if entry and entry["num"] is not None else (next(iter(nums)) if len(nums) == 1 else None)
    if hit := _by_number(num, images, _nums):
        return hit
    if entry:  # 썸네일 이름이 제목으로 돼 있는 경우
        score = lambda i: max(_sim(i.stem, t) for t in [entry["orig"], entry["title"], *entry["titles"]])  # noqa: E731
        best = max(images, key=score, default=None)
        if best and score(best) >= 0.6:
            return best
    return _find_pair(video, images, videos)


def scan_channel(folder):
    """[(영상, 메모장 항목 또는 None, 썸네일 또는 None)] — 한 묶음의 한 채널 폴더 안에서, 파일 이름 순서."""
    files = [p for p in Path(folder).rglob("*") if p.is_file()]
    videos = sorted((p for p in files if p.suffix.lower() in VIDEO_EXT), key=natural_key)
    entries = load_entries(p for p in files if p.suffix.lower() == ".txt")
    images = [p for p in files if p.suffix.lower() in IMAGE_EXT]
    pairs = []
    for v in videos:
        entry = _match_entry(v, entries, videos)
        pairs.append([v, entry, _match_thumb(v, entry, images, videos)])
    for col in (1, 2):  # 한 메모장 항목·썸네일이 두 영상에 짝지어지면 둘 다 확인 필요로 돌린다
        used = [id(p[col]) for p in pairs if p[col]]
        for p in pairs:
            if p[col] and used.count(id(p[col])) > 1:
                p[col] = None
    return [tuple(p) for p in pairs]


def channel_dirs(ch):
    """영상/<묶음>/<채널 폴더> 들, 묶음 이름 순서."""
    return sorted((d for d in VIDEO_DIR.glob(f"*/{ch['folder']}") if d.is_dir()), key=natural_key)


def new_batch(cfg, today=None):
    """영상/<오늘 날짜>/<채널 폴더들> 을 만든다. 같은 날 또 만들면 -2, -3 …"""
    name = (today or datetime.now(KST)).strftime("%Y-%m-%d")
    batch, n = VIDEO_DIR / name, 2
    while batch.exists():
        batch, n = VIDEO_DIR / f"{name}-{n}", n + 1
    for ch in cfg["channels"]:
        (batch / ch["folder"]).mkdir(parents=True)
    return batch


# ---------- 예약 ----------
DAYS = "월화수목금토일"


def parse_slots(spec):
    out = []
    for part in re.split(r"[,\n/]+", spec or ""):
        if not part.strip():
            continue
        m = re.fullmatch(r"\s*([월화수목금토일])\S*\s+(\d{1,2}):(\d{2})\s*", part)
        if not m or int(m[2]) > 23 or int(m[3]) > 59:
            raise ValueError(f"예약 시간 형식 오류: '{part.strip()}' (예: 월 18:00)")
        out.append((DAYS.index(m[1]), int(m[2]), int(m[3])))
    if not out:
        raise ValueError("예약 요일·시간이 비어 있습니다.")
    return sorted(out)


def next_slots(after, slots, n):
    res, d = [], after.date()
    while len(res) < n:
        for wd, h, mi in slots:
            t = datetime(d.year, d.month, d.day, h, mi, tzinfo=KST)
            if wd == d.weekday() and t > after and len(res) < n:
                res.append(t)
        d += timedelta(days=1)
    return res


def parse_dt(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=KST)
    return datetime.strptime(str(value).strip()[:16], "%Y-%m-%d %H:%M").replace(tzinfo=KST)


def fmt_dt(t):
    return t.astimezone(KST).strftime("%Y-%m-%d %H:%M")


def fill_schedule(rows, ch, last_scheduled, now):
    """예약일시가 빈 행을 채널의 다음 빈 자리부터 채운다."""
    mine = [r for r in rows if r["채널"] == ch["name"]]
    blanks = [r for r in mine if not r["예약일시"]]
    if not blanks:
        return
    after = max([now + timedelta(minutes=30)] + ([last_scheduled] if last_scheduled else [])
                + [parse_dt(r["예약일시"]) for r in mine if r["예약일시"]])
    for r, t in zip(blanks, next_slots(after, parse_slots(ch["slots"]), len(blanks))):
        r["예약일시"] = fmt_dt(t)


# ---------- 엑셀 ----------
def rel(p):
    return Path(p).relative_to(VIDEO_DIR).as_posix() if p else ""


def build_rows(cfg, old_rows):
    """영상 폴더를 훑어 엑셀 행을 만든다. 이미 있던 행(사용자가 고친 제목 등)은 그대로 둔다."""
    old = {r["파일"]: r for r in old_rows}
    rows, problems = [], []
    for ch, folder in ((ch, d) for ch in cfg["channels"] for d in channel_dirs(ch)):
        added = []
        for video, entry, thumb in scan_channel(folder):
            name = rel(video)
            if name in old:
                rows.append(old[name])
                continue
            info = entry or {}
            if not entry:
                problems.append(f"{name}: 메모장(SEO) 내용을 못 찾음")
            if not thumb:
                problems.append(f"{name}: 썸네일을 못 찾음")
            elif thumb.stat().st_size > THUMB_MAX:
                problems.append(f"{name}: 썸네일이 2MB를 넘음 ({thumb.name})")
            row = dict.fromkeys(COLUMNS, "") | {
                "채널": ch["name"], "파일": name, "제목": info.get("title", ""),
                "제목후보": "\n".join(f"{i}. {t}" for i, t in enumerate(info.get("titles", []), 1)),
                "설명": info.get("desc", ""), "태그": info.get("tags", ""), "고정댓글": info.get("comment", ""),
                "메모장": rel(info.get("src")), "썸네일": rel(thumb)}
            rows.append(row)
            added.append(row)
        if ch["members"] and added and not any(r["회원우선"] for r in rows if r["채널"] == ch["name"]):
            added[-1]["회원우선"] = "O"  # 기본 규칙: 이번 묶음의 마지막 영상
    names = {ch["name"] for ch in cfg["channels"]}
    problems += [f"{r['파일']}: 설정에 없는 채널 '{r['채널']}'" for r in rows if r["채널"] not in names]
    unknown = {rel(v) for v in VIDEO_DIR.rglob("*") if v.suffix.lower() in VIDEO_EXT} - {r["파일"] for r in rows}
    problems += [f"{name}: 어느 채널인지 모름 (영상/묶음/채널폴더/ 안에 넣어주세요)" for name in sorted(unknown)]
    return rows, problems


def read_rows():
    if not EXCEL.exists():
        return []
    import openpyxl
    ws = openpyxl.load_workbook(EXCEL).active
    head = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        r = {col: values[head.index(col)] if col in head else "" for col in COLUMNS}
        r = {k: ("" if v is None else fmt_dt(parse_dt(v)) if isinstance(v, datetime) else str(v))
             for k, v in r.items()}
        if r["파일"]:
            rows.append(r)
    return rows


def write_rows(rows):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "업로드목록"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
    widths = {"채널": 12, "파일": 30, "제목": 45, "제목후보": 45, "설명": 40, "태그": 25, "고정댓글": 30, "메모장": 20,
              "썸네일": 20, "회원우선": 9, "예약일시": 17, "상태": 12, "오류": 30, "영상ID": 14}
    edit = PatternFill("solid", fgColor="FFF2CC")  # 사용자가 고치는 칸
    for i, col in enumerate(COLUMNS, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths[col]
        ws.cell(1, i).font = Font(bold=True)
        for cell in ws[letter][1:]:
            cell.alignment = Alignment(wrap_text=col in ("제목후보", "설명", "고정댓글", "오류"), vertical="top")
            if col in ("제목", "회원우선", "예약일시"):
                cell.fill = edit
    ws.freeze_panes = "C2"
    try:
        wb.save(EXCEL)
    except PermissionError:
        raise PermissionError("엑셀 파일이 열려 있어서 저장하지 못했습니다. 엑셀을 닫고 다시 눌러주세요.") from None


# ---------- 완료 ----------
def _move(relpath):
    src = VIDEO_DIR / relpath
    if not relpath or not src.exists():
        return
    dst = DONE_DIR / relpath
    n = 1
    while dst.exists():  # 같은 이름이 이미 있으면 덮어쓰지 않는다
        dst = dst.with_stem(f"{Path(relpath).stem} ({n})")
        n += 1
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(src, dst)


def move_done(row):
    """영상·썸네일을 완료/로. 이 묶음의 채널 폴더에 영상이 안 남으면 나머지(메모장·대본 등)도 통째로."""
    _move(row["파일"])
    _move(row.get("썸네일", ""))
    batch = VIDEO_DIR / Path(row["파일"]).parts[0]
    unit = batch / Path(row["파일"]).parts[1]
    rest = list(unit.rglob("*"))
    if any(p.suffix.lower() in VIDEO_EXT for p in rest):
        return
    for f in (p for p in rest if p.is_file()):
        _move(rel(f))
    dirs = sorted((p for p in rest if p.is_dir()), key=lambda p: len(p.parts), reverse=True)
    for d in dirs + [unit, batch]:
        if d.exists() and not any(d.iterdir()):
            d.rmdir()
