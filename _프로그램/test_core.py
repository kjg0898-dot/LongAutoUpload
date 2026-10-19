"""YouTube 연결 없이 확인할 수 있는 부분 검사. 실행: python test_core.py"""
import tempfile
from datetime import datetime
from pathlib import Path

import core
from core import KST


def test_merge():
    m = core.merge
    assert m("【명언】 | 채널", "인생 명언", "기준 글자 앞", "|") == "【명언】 인생 명언 | 채널"
    assert m("【명언】 | 채널", "인생 명언", "기준 글자 뒤", "【명언】") == "【명언】 인생 명언 | 채널"
    assert m("A {제목} B", "X", "표시 문자 바꾸기", "{제목}") == "A X B"
    assert m("구독 부탁", "본문", "맨 앞", sep="\n\n") == "본문\n\n구독 부탁"
    assert m("구독 부탁", "본문", "맨 뒤", sep="\n\n") == "구독 부탁\n\n본문"
    # 제목이 틀에 있는 낱말(예: 고전)과 같아도 넣는다 (다시 실행 대비는 원본 틀 백업으로 함)
    assert m("[숨겨진 동양사]   l 역사 인물 l 고전", "고전", "기준 글자 뒤", "]") == "[숨겨진 동양사] 고전 l 역사 인물 l 고전"
    assert m("", "제목만", "기준 글자 앞", "|") == "제목만"
    assert m("틀", "", "맨 앞") == "틀"
    try:
        m("틀에 기준 없음", "X", "기준 글자 앞", "|")
        raise AssertionError
    except ValueError:
        pass


def test_tags():
    assert core.parse_tags("#명언, 좋은글\n#인생 #명언 모음") == ["명언", "좋은글", "인생", "명언 모음"]
    assert core.merge_tags(["명언", "기본"], ["명언", "새것"], "뒤에 추가") == ["명언", "기본", "새것"]
    assert core.merge_tags(["기본"], ["새것"], "앞에 추가") == ["새것", "기본"]
    assert core.tags_len(["a b", "c"]) == 3 + 2 + 1 + 1
    many = core.merge_tags([], [f"태그{i:03}" for i in range(200)], "뒤에 추가")
    assert core.tags_len(many) <= 500 < core.tags_len(many + ["x" * 6])


def test_memo():
    text = "제목: 인생 명언 10가지\n\n[설명]\n첫 줄\n둘째 줄\n\n■ 태그 : #명언, #인생\n고정 댓글:\n여러분의 명언은?\n"
    assert core.parse_memo(text) == {"title": "인생 명언 10가지", "desc": "첫 줄\n둘째 줄",
                                     "tags": "#명언, #인생", "comment": "여러분의 명언은?"}
    assert core.parse_memo("머리말 없는 글") == {}
    assert core.parse_memo("설명하자면 이렇다") == {}  # 낱말이 머리말로 오인되지 않게


SEO = """============================================================
[2026-10-17 21:19] 행 2 | 순번 1 | 제목: 원본 제목 하나
============================================================

▣ 채널 판별: [역사]
>> 원본 제목: [원본 제목 하나]
---
제목 (3안)
1. 첫째 후보
2. 둘째 후보
3. 셋째 후보
추천 1순위: 2 — 이유
---
썸네일 문구 (3안)
1. 문구
추천 1순위: 1 — 이유
---
설명글 & 해시태그
설명 본문입니다.

참고자료
· 책

#역사 #동양사
---
고정 댓글
고마워요 ❤️
둘째 줄
---
영상 프롬프트
[시퀀스 1]
============================================================
[2026-10-17 21:35] 행 3 | 순번 2 | 제목: 원본 둘
============================================================
---
제목 (3안)
1. 두 번째 영상 제목
---
설명글 & 해시태그
둘째 설명
---
고정 댓글
둘째 댓글
"""


def test_seo():
    a, b = core.parse_seo(SEO.replace("\n", "\r\n"))
    assert a == {"num": 1, "orig": "원본 제목 하나", "titles": ["첫째 후보", "둘째 후보", "셋째 후보"],
                 "title": "둘째 후보", "desc": "설명 본문입니다.\n\n참고자료\n· 책\n\n#역사 #동양사",
                 "tags": "", "comment": "고마워요 ❤️\n둘째 줄"}, a
    assert (b["num"], b["title"], b["comment"], b["tags"]) == (2, "두 번째 영상 제목", "둘째 댓글", "")
    assert core.parse_seo("제목: 일반 메모장") == []


def test_seo_batch_folder():
    """폴더 모양: 영상/<묶음>/<채널>/{seo,썸네일,대본,...} — 번호로 짝짓고, 끝나면 채널 폴더째 완료로."""
    with tempfile.TemporaryDirectory() as d:
        core.VIDEO_DIR, core.DONE_DIR = Path(d) / "영상", Path(d) / "완료"
        cfg = {"channels": [{**core.DEFAULT_CHANNEL, "name": "동양사", "folder": "동양사"},
                            {**core.DEFAULT_CHANNEL, "name": "고전", "folder": "고전"}]}
        batch = core.new_batch(cfg, datetime(2026, 10, 19))
        assert batch.name == "2026-10-19" and core.new_batch(cfg, datetime(2026, 10, 19)).name == "2026-10-19-2"
        b = batch / "동양사"
        for sub in ("seo", "썸네일", "대본"):
            (b / sub).mkdir()
        (b / "seo" / "20261017_seo.txt").write_text(SEO, "utf-8")
        (b / "대본" / "1.txt").write_text("대본은 메모장이 아님", "utf-8")
        for n in (1, 2):
            (b / f"숨겨진동양사이야기_{n}_내보내기.mp4").write_bytes(b"v")
            (b / "썸네일" / f"썸네일{n}.png").write_bytes(b"i")
        (batch / "고전" / "고전_1.mp4").write_bytes(b"v")  # 같은 묶음의 다른 채널: 번호가 같아도 섞이면 안 됨
        (core.VIDEO_DIR / "2026-10-19-2" / "어디채널").mkdir()
        (core.VIDEO_DIR / "2026-10-19-2" / "어디채널" / "x.mp4").write_bytes(b"v")
        rows, problems = core.build_rows(cfg, [])
        new = [r for r in rows if r["채널"] == "동양사"]
        assert [(r["제목"], r["썸네일"].split("/")[-1], r["고정댓글"], r["태그"]) for r in new] == [
            ("둘째 후보", "썸네일1.png", "고마워요 ❤️\n둘째 줄", ""), ("두 번째 영상 제목", "썸네일2.png", "둘째 댓글", "")], new
        assert new[0]["제목후보"] == "1. 첫째 후보\n2. 둘째 후보\n3. 셋째 후보"
        assert [r["파일"] for r in rows if r["채널"] == "고전"] == ["2026-10-19/고전/고전_1.mp4"]
        assert sorted(p.split(":")[0] for p in problems) == [
            "2026-10-19-2/어디채널/x.mp4", "2026-10-19/고전/고전_1.mp4", "2026-10-19/고전/고전_1.mp4"], problems

        core.move_done(new[0])
        assert (b / "seo" / "20261017_seo.txt").exists()  # 영상이 남아 있으면 메모장은 그대로
        core.move_done(new[1])
        assert not b.exists() and (core.DONE_DIR / "2026-10-19" / "동양사" / "대본" / "1.txt").exists()
        assert len(list((core.DONE_DIR / "2026-10-19" / "동양사").rglob("*.*"))) == 6
        assert (batch / "고전").exists()  # 다른 채널 폴더는 그대로


def test_slots():
    slots = core.parse_slots("월 18:00, 수요일 18:00, 금 20:30")
    assert slots == [(0, 18, 0), (2, 18, 0), (4, 20, 30)]
    after = datetime(2026, 10, 19, 18, 0, tzinfo=KST)  # 월요일 18:00 정각 (이 자리는 제외)
    got = [core.fmt_dt(t) for t in core.next_slots(after, slots, 4)]
    assert got == ["2026-10-21 18:00", "2026-10-23 20:30", "2026-10-26 18:00", "2026-10-28 18:00"], got
    for bad in ("", "월18:00", "월 25:00", "x 18:00"):
        try:
            core.parse_slots(bad)
            raise AssertionError(bad)
        except ValueError:
            pass


def test_fill_schedule():
    ch = {"name": "A", "slots": "화 09:00"}
    rows = [{"채널": "A", "예약일시": "2026-11-03 09:00"}, {"채널": "A", "예약일시": ""},
            {"채널": "B", "예약일시": ""}]
    now = datetime(2026, 10, 16, 12, 0, tzinfo=KST)
    core.fill_schedule(rows, ch, datetime(2026, 10, 27, 9, 0, tzinfo=KST), now)
    assert rows[1]["예약일시"] == "2026-11-10 09:00"  # 엑셀에 이미 잡힌 마지막 자리 다음
    assert rows[2]["예약일시"] == ""


def test_scan_and_build():
    with tempfile.TemporaryDirectory() as d:
        core.VIDEO_DIR = Path(d) / "영상"
        a = core.VIDEO_DIR / "묶음" / "명언채널"
        (a / "02_사랑").mkdir(parents=True)
        for name in ("10_끝.mp4", "10_끝.txt", "10_끝.jpg", "1_시작.mp4", "1_시작 메모.txt", "1_시작_썸네일.png"):
            (a / name).write_bytes(b"x")
        (a / "1_시작 메모.txt").write_text("제목: 시작\n설명: 첫 영상\n태그: 명언", "utf-8")
        (a / "10_끝.txt").write_bytes("제목: 끝".encode("cp949"))
        for name in ("영상.mp4", "info.txt", "cover.jpg"):  # 영상마다 폴더가 따로 있는 경우
            (a / "02_사랑" / name).write_bytes(b"x")
        pairs = core.scan_channel(a)
        assert [(v.name, m and m["src"].name, t and t.name) for v, m, t in pairs] == [
            ("1_시작.mp4", "1_시작 메모.txt", "1_시작_썸네일.png"),
            ("영상.mp4", None, "cover.jpg"),  # 내용 없는 메모장은 짝으로 치지 않는다
            ("10_끝.mp4", "10_끝.txt", "10_끝.jpg")], pairs

        cfg = {"channels": [{**core.DEFAULT_CHANNEL, "name": "명언", "folder": "명언채널", "members": True}]}
        rows, problems = core.build_rows(cfg, [])
        assert [r["회원우선"] for r in rows] == ["", "", "O"]
        assert rows[0]["제목"] == "시작" and rows[0]["태그"] == "명언" and rows[2]["제목"] == "끝"
        assert any("영상.mp4" in p and "메모장" in p for p in problems)  # 빈 메모장은 확인 필요로 알린다
        rows[0]["제목"] = "고친 제목"
        again, _ = core.build_rows(cfg, rows)
        assert again[0]["제목"] == "고친 제목"  # 다시 만들어도 사용자가 고친 칸은 그대로

        core.EXCEL = Path(d) / "목록.xlsx"
        core.write_rows(rows)
        assert core.read_rows() == rows

        core.DONE_DIR = Path(d) / "완료"
        core.move_done(rows[2])
        assert (core.DONE_DIR / rows[2]["파일"]).exists() and not (core.VIDEO_DIR / rows[2]["파일"]).exists()


def test_excel_safety():
    """다른 시트는 남고, 날짜만 적은 예약일시·이상한 값은 형식 오류로, 회원우선은 O만 예."""
    import openpyxl
    with tempfile.TemporaryDirectory() as d:
        core.EXCEL = Path(d) / "목록.xlsx"
        row = dict.fromkeys(core.COLUMNS, "") | {"채널": "A", "파일": "b/A/1.mp4", "제목": "제목"}
        core.write_rows([row])
        wb = openpyxl.load_workbook(core.EXCEL)
        wb.create_sheet("메모")["A1"] = "내 메모"
        ws = wb[core.SHEET]
        ws.cell(2, core.COLUMNS.index("예약일시") + 1).value = datetime(2026, 10, 21)  # 엑셀에서 날짜만 입력
        wb.active = 1  # 사용자가 다른 시트를 보고 저장
        wb.save(core.EXCEL)
        rows = core.read_rows()
        assert rows[0]["제목"] == "제목" and rows[0]["예약일시"] == "2026-10-21 00:00"  # 4단계에서 공개 시간과 다르다고 확인
        for bad in ("11월 3일 18시", "2026-10-21"):
            try:
                core.parse_dt(bad)
                raise AssertionError(bad)
            except ValueError as e:
                assert "형식" in str(e)
        core.write_rows(rows)
        assert openpyxl.load_workbook(core.EXCEL)["메모"]["A1"].value == "내 메모"
        assert not list(Path(d).glob("~*"))  # 임시 파일이 남지 않음
    assert [core.is_member(v) for v in ("O", "o", "예", "", "X", "아니요")] == [True, True, True, False, False, False]


def test_config_validation():
    with tempfile.TemporaryDirectory() as d:
        core.CONFIG, core.VIDEO_DIR = Path(d) / "설정.json", Path(d) / "영상"
        ch = lambda name, folder, slots="월 12:00": {**core.DEFAULT_CHANNEL, "name": name, "folder": folder, "slots": slots}  # noqa: E731
        for bad, word in (([ch("A", "x"), ch("B", "X")], "폴더"), ([ch("A", "a/b")], "글자"),
                          ([ch("A", "")], "폴더"), ([ch("숨겨진동양사", "a", "월요일12시")], "[숨겨진동양사]")):
            try:
                core.save_config({**core.DEFAULT_CONFIG, "channels": bad})
                raise AssertionError(bad)
            except ValueError as e:
                assert word in str(e), e
        core.save_config({**core.DEFAULT_CONFIG, "channels": [ch("A", "[숨겨진동양사]")]})
        (core.VIDEO_DIR / "1" / "[숨겨진동양사]").mkdir(parents=True)
        assert len(core.channel_dirs(ch("A", "[숨겨진동양사]"))) == 1  # [ ] 가 들어간 폴더 이름


def test_overrides():
    ch = {**core.DEFAULT_CHANNEL}
    assert core.overrides(ch) == ({}, {})  # 기본값은 아무것도 건드리지 않음
    ch.update(category="교육", synthetic="아니요", license="크리에이티브 커먼즈", language="한국어")
    assert core.overrides(ch) == ({"categoryId": "27", "defaultLanguage": "ko", "defaultAudioLanguage": "ko"},
                                  {"containsSyntheticMedia": False, "license": "creativeCommon"})


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
